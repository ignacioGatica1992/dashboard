import streamlit as st
import pandas as pd
import io

st.set_page_config(
    page_title="Doc → Markdown Converter",
    page_icon="📄",
    layout="wide",
)

st.title("📄 Document to Markdown Converter")
st.caption("Convert Word, Excel, PowerPoint and PDF files to Markdown format")

SUPPORTED_TYPES = {
    "docx": "Word (.docx)",
    "xlsx": "Excel (.xlsx)",
    "xls": "Excel (.xls)",
    "pptx": "PowerPoint (.pptx)",
    "pdf": "PDF (.pdf)",
}


def word_to_markdown(file_bytes: bytes) -> str:
    import mammoth
    result = mammoth.convert_to_markdown(io.BytesIO(file_bytes))
    messages = result.messages
    md = result.value
    if messages:
        warnings = "\n".join(f"> ⚠️ {m.message}" for m in messages)
        md = f"{warnings}\n\n{md}"
    return md


def excel_to_markdown(file_bytes: bytes) -> str:
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    parts = []
    for sheet_name in xl.sheet_names:
        df = xl.parse(sheet_name)
        parts.append(f"## {sheet_name}\n")
        parts.append(df.to_markdown(index=False))
        parts.append("\n")
    return "\n".join(parts)


def pptx_to_markdown(file_bytes: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"## Slide {i}\n")
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            is_title = (
                hasattr(shape, "placeholder_format")
                and shape.placeholder_format is not None
                and shape.placeholder_format.idx == 0
            )
            if is_title:
                title_text = shape.text.strip()
                if title_text:
                    parts.append(f"### {title_text}\n")
            else:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    indent = "  " * para.level
                    parts.append(f"{indent}- {text}")
        parts.append("\n---\n")
    return "\n".join(parts)


def pdf_to_markdown(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF

    parts = []
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    for i, page in enumerate(doc, start=1):
        parts.append(f"## Page {i}\n")
        text = page.get_text("text")
        if text.strip():
            parts.append(text.strip())
        parts.append("\n")
    doc.close()
    return "\n".join(parts)


CONVERTERS = {
    "docx": word_to_markdown,
    "xlsx": excel_to_markdown,
    "xls": excel_to_markdown,
    "pptx": pptx_to_markdown,
    "pdf": pdf_to_markdown,
}


uploaded_file = st.file_uploader(
    "Upload your document",
    type=list(SUPPORTED_TYPES.keys()),
    help="Supported: Word (.docx), Excel (.xlsx/.xls), PowerPoint (.pptx), PDF (.pdf)",
)

if uploaded_file is not None:
    ext = uploaded_file.name.rsplit(".", 1)[-1].lower()
    file_bytes = uploaded_file.read()

    st.info(f"**File:** {uploaded_file.name}  |  **Type:** {SUPPORTED_TYPES.get(ext, ext.upper())}  |  **Size:** {len(file_bytes)/1024:.1f} KB")

    with st.spinner("Converting…"):
        try:
            converter = CONVERTERS[ext]
            markdown_output = converter(file_bytes)
            success = True
        except Exception as exc:
            st.error(f"Conversion failed: {exc}")
            success = False

    if success:
        col1, col2 = st.columns(2)

        with col1:
            st.subheader("Markdown source")
            st.code(markdown_output, language="markdown")

        with col2:
            st.subheader("Preview")
            st.markdown(markdown_output)

        st.divider()

        output_filename = uploaded_file.name.rsplit(".", 1)[0] + ".md"
        st.download_button(
            label="⬇️ Download Markdown file",
            data=markdown_output.encode("utf-8"),
            file_name=output_filename,
            mime="text/markdown",
        )
